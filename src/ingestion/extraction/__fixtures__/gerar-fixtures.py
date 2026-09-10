#!/usr/bin/env python3
"""Regera os arquivos binarios de fixture usados pelos testes de extracao.

Os binarios ficam versionados no repositorio: e o que permite a suite rodar em
CI sem Office, LibreOffice ou macOS instalados. Este script existe para quando
o conteudo esperado mudar, e nao como etapa de build.

Uso (macOS, por causa do textutil que gera o .doc):

    python3 -m venv /tmp/fixgen
    /tmp/fixgen/bin/pip install xlwt python-pptx
    /tmp/fixgen/bin/python gerar-fixtures.py

Os valores literais daqui sao os mesmos que os testes procuram. Mudar um texto
aqui quebra o teste correspondente de proposito: a fixture e a fonte da verdade
do que a extracao precisa preservar.
"""

import datetime
import os
import struct
import subprocess
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))

DOC_HTML = """<html><body>
<p>Manual de marca legado da Acme Corporation</p>
<p>O codigo de identificacao da marca e ORQUIDEA CROMADA 47.</p>
<table border="1">
<tr><td>Cor primaria</td><td>azul-cobalto</td></tr>
<tr><td>Cor secundaria</td><td>areia-quente</td></tr>
</table>
<p>Aplicacoes aprovadas:</p>
<ul><li>Fachada da loja</li><li>Frota de entrega</li></ul>
</body></html>
"""

SOMA_ESPERADA = 23250.5


def gerar_doc():
    """Word 97-2003 binario, com paragrafo, tabela e lista."""
    destino = os.path.join(AQUI, "legado.doc")
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False) as origem:
        origem.write(DOC_HTML)
        caminho_origem = origem.name
    try:
        subprocess.run(
            ["textutil", "-convert", "doc", "-output", destino, caminho_origem],
            check=True,
        )
    finally:
        os.unlink(caminho_origem)
    print("gerado", destino)


def _gravar_resultado_de_formula(caminho, esperado):
    """Escreve o resultado em cache da formula, que o xlwt deixa em branco.

    O Excel grava o ultimo valor calculado dentro do proprio registro FORMULA, e
    e esse valor que a extracao le -- ninguem recalcula a planilha na ingestao.
    O xlwt grava o tipo especial "string vazia" no lugar, entao a fixture nao
    exercitaria o caminho que importa sem esta correcao.
    """
    with open(caminho, "rb") as arquivo:
        dados = bytearray(arquivo.read())

    cabecalho = struct.pack("<HHHHH", 0x0006, 0x0023, 3, 1, 17)
    marcador_em_branco = bytes([0x03, 0, 0, 0, 0, 0, 0xFF, 0xFF])
    alvo = cabecalho + marcador_em_branco

    posicao = dados.find(alvo)
    if posicao < 0:
        sys.exit("registro FORMULA nao encontrado; o layout do xlwt mudou")

    inicio = posicao + len(cabecalho)
    dados[inicio:inicio + 8] = struct.pack("<d", esperado)

    with open(caminho, "wb") as arquivo:
        arquivo.write(bytes(dados))


def gerar_xls():
    """Excel 97-2003 binario: duas abas, datas, numeros, formula e buraco de linha."""
    import xlwt

    destino = os.path.join(AQUI, "legado.xls")
    livro = xlwt.Workbook(encoding="utf-8")
    estilo_data = xlwt.XFStyle()
    estilo_data.num_format_str = "DD/MM/YYYY"

    midia = livro.add_sheet("Midia")
    for coluna, titulo in enumerate(["Canal", "Verba", "Inicio"]):
        midia.write(0, coluna, titulo)
    midia.write(1, 0, "Instagram")
    midia.write(1, 1, 15000)
    midia.write(1, 2, datetime.date(2026, 10, 1), estilo_data)
    midia.write(2, 0, "Radio")
    midia.write(2, 1, 8250.5)
    midia.write(2, 2, datetime.date(2026, 11, 15), estilo_data)
    midia.write(3, 0, "Total")
    midia.write(3, 1, xlwt.Formula("SUM(B2:B3)"))

    cronograma = livro.add_sheet("Cronograma")
    cronograma.write(0, 0, "Etapa")
    cronograma.write(0, 1, "Prazo")
    cronograma.write(1, 0, "Entrega ORQUIDEA CROMADA 47")
    cronograma.write(1, 1, datetime.date(2026, 12, 20), estilo_data)
    cronograma.write(3, 0, "Observacao depois de uma linha vazia")
    cronograma.write(3, 2, "coluna C com a B vazia")

    livro.save(destino)
    _gravar_resultado_de_formula(destino, SOMA_ESPERADA)
    print("gerado", destino)


def gerar_pptx():
    """Apresentacao com texto, tabela, notas e um slide feito apenas de imagem."""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    destino = os.path.join(AQUI, "apresentacao.pptx")
    apresentacao = Presentation()

    titulo_e_conteudo = apresentacao.slide_layouts[1]
    slide = apresentacao.slides.add_slide(titulo_e_conteudo)
    slide.shapes.title.text = "Campanha Verao 2026"
    corpo = slide.placeholders[1].text_frame
    corpo.text = "Praca principal: litoral norte"
    corpo.add_paragraph().text = "Verba aprovada: 480 mil"
    slide.notes_slide.notes_text_frame.text = (
        "Nota do apresentador: confirmar ORQUIDEA CROMADA 47 com o cliente."
    )

    somente_titulo = apresentacao.slide_layouts[5]
    slide_tabela = apresentacao.slides.add_slide(somente_titulo)
    slide_tabela.shapes.title.text = "Divisao por canal"
    tabela = slide_tabela.shapes.add_table(
        3, 2, Inches(1), Inches(2), Inches(6), Inches(2)
    ).table
    tabela.cell(0, 0).text = "Canal"
    tabela.cell(0, 1).text = "Participacao"
    tabela.cell(1, 0).text = "Digital"
    tabela.cell(1, 1).text = "62 por cento"
    tabela.cell(2, 0).text = "Radio"
    tabela.cell(2, 1).text = "38 por cento"

    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    slide_grafico = apresentacao.slides.add_slide(somente_titulo)
    slide_grafico.shapes.title.text = "Investimento por praca"
    dados = CategoryChartData()
    dados.categories = ["Litoral Norte", "Serra Gaucha"]
    dados.add_series("Verba 2026", (280.0, 200.0))
    slide_grafico.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), dados
    )

    em_branco = apresentacao.slide_layouts[6]
    slide_imagem = apresentacao.slides.add_slide(em_branco)
    slide_imagem.shapes.add_picture(
        _png_solido(), Inches(1), Inches(1), Inches(4), Inches(3)
    )

    apresentacao.save(destino)
    print("gerado", destino)


def _png_solido():
    """Um PNG minimo e valido, para o slide que so tem imagem."""
    import io
    import zlib

    largura = altura = 8
    linhas = b"".join(b"\x00" + bytes([200, 60, 40] * largura) for _ in range(altura))

    def parte(tipo, dados):
        bruto = tipo + dados
        return struct.pack(">I", len(dados)) + bruto + struct.pack(
            ">I", zlib.crc32(bruto) & 0xFFFFFFFF
        )

    cabecalho = struct.pack(">IIBBBBB", largura, altura, 8, 2, 0, 0, 0)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + parte(b"IHDR", cabecalho)
        + parte(b"IDAT", zlib.compress(linhas))
        + parte(b"IEND", b"")
    )
    return io.BytesIO(png)


PDF_LINHAS = [
    "Guia de marca da Acme Corporation",
    "O codigo de identificacao da marca e ORQUIDEA CROMADA 47.",
    "A cor primaria e o azul profundo #12385A.",
    "E proibido aplicar o logotipo sobre fundo vermelho.",
]


def gerar_pdf():
    """Escreve um PDF de uma pagina com camada de texto, sem dependencia externa.

    O arquivo existe para o teste em tela ter um PDF de verdade para enviar pelo
    seletor de arquivos, e para a suite de formatos exercer o extrator de PDF
    pelo mesmo caminho dos outros formatos.
    """

    conteudo = "BT /F1 12 Tf 72 760 Td 16 TL\n"
    for linha in PDF_LINHAS:
        conteudo += f"({linha}) Tj T*\n"
    conteudo += "ET"

    objetos = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(conteudo)} >>\nstream\n{conteudo}\nendstream",
    ]

    pdf = "%PDF-1.4\n"
    deslocamentos = []
    for numero, objeto in enumerate(objetos, start=1):
        deslocamentos.append(len(pdf))
        pdf += f"{numero} 0 obj\n{objeto}\nendobj\n"

    inicio_xref = len(pdf)
    pdf += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n"
    for deslocamento in deslocamentos:
        pdf += f"{deslocamento:010d} 00000 n \n"
    pdf += (
        f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
        f"startxref\n{inicio_xref}\n%%EOF\n"
    )

    destino = os.path.join(AQUI, "guia-marca.pdf")
    with open(destino, "wb") as arquivo:
        arquivo.write(pdf.encode("latin-1"))
    print(f"gerado: {destino}")


if __name__ == "__main__":
    gerar_doc()
    gerar_xls()
    gerar_pptx()
    gerar_pdf()
